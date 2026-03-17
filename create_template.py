#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""创建示例Word模板文件"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_template():
    # 创建文档
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    
    # 标题
    title = doc.add_heading('行政处罚决定书', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 文号
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run('案号：{{案号}}')
    
    doc.add_paragraph()  # 空行
    
    # 当事人信息
    doc.add_heading('一、当事人信息', level=1)
    
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Table Grid'
    
    # 填充表格
    cells = [
        ('当事人名称', '{{当事人名称}}'),
        ('个人姓名', '{{个人姓名}}'),
        ('性别', '{{性别}}'),
        ('身份证号', '{{身份证号}}'),
        ('联系电话', '{{个人电话}}'),
        ('家庭地址', '{{家庭地址}}'),
    ]
    
    for i, (label, value) in enumerate(cells):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = value
    
    doc.add_paragraph()
    
    # 单位信息
    doc.add_heading('二、单位信息', level=1)
    
    table2 = doc.add_table(rows=4, cols=2)
    table2.style = 'Table Grid'
    
    cells2 = [
        ('单位名称', '{{单位名称}}'),
        ('统一社会信用代码', '{{统一社会信用代码}}'),
        ('法人姓名', '{{法人姓名}}'),
        ('法人职务', '{{法人职务}}'),
    ]
    
    for i, (label, value) in enumerate(cells2):
        row = table2.rows[i]
        row.cells[0].text = label
        row.cells[1].text = value
    
    doc.add_paragraph()
    
    # 案件信息
    doc.add_heading('三、案件信息', level=1)
    
    p = doc.add_paragraph()
    p.add_run('案件名称：').bold = True
    p.add_run('{{案件名称}}')
    
    p = doc.add_paragraph()
    p.add_run('案发地点：').bold = True
    p.add_run('{{案发地点}}')
    
    p = doc.add_paragraph()
    p.add_run('案件来源：').bold = True
    p.add_run('{{案件来源}}')
    
    p = doc.add_paragraph()
    p.add_run('案由：').bold = True
    p.add_run('{{案由}}')
    
    doc.add_paragraph()
    
    # 违法事实
    doc.add_heading('四、违法事实', level=1)
    doc.add_paragraph('{{违法事实}}')
    
    doc.add_paragraph()
    
    # 处罚依据
    doc.add_heading('五、处罚依据及决定', level=1)
    
    p = doc.add_paragraph()
    p.add_run('违法条款：').bold = True
    p.add_run('{{违法条款}}')
    
    p = doc.add_paragraph()
    p.add_run('处罚意见：').bold = True
    p.add_run('{{处罚意见}}')
    
    p = doc.add_paragraph()
    p.add_run('处罚决定：').bold = True
    p.add_run('{{处罚决定}}')
    
    doc.add_paragraph()
    
    # 承办人信息
    doc.add_heading('六、承办人信息', level=1)
    
    p = doc.add_paragraph()
    p.add_run('承办人：').bold = True
    p.add_run('{{承办人}}')
    
    p = doc.add_paragraph()
    p.add_run('承办人1：').bold = True
    p.add_run('{{承办人姓名1}}（执法证号：{{执法证号1}}）')
    
    p = doc.add_paragraph()
    p.add_run('承办人2：').bold = True
    p.add_run('{{承办人姓名2}}（执法证号：{{执法证号2}}）')
    
    doc.add_paragraph()
    
    # 项目信息
    doc.add_heading('七、项目信息', level=1)
    
    table3 = doc.add_table(rows=5, cols=2)
    table3.style = 'Table Grid'
    
    cells3 = [
        ('项目名称', '{{项目名称}}'),
        ('建设单位', '{{建设单位}}'),
        ('施工单位', '{{施工单位}}'),
        ('监理单位', '{{监理单位}}'),
        ('建设规模', '{{建设规模}}'),
    ]
    
    for i, (label, value) in enumerate(cells3):
        row = table3.rows[i]
        row.cells[0].text = label
        row.cells[1].text = value
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # 落款
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run('生成日期：{{生成日期}}')
    
    # 保存
    output_path = "G:/XZCF/行政处罚生成工具/模板文件/行政处罚.docx"
    doc.save(output_path)
    print(f"模板文件已创建: {output_path}")
    print("\n模板中包含的占位符：")
    print("- {{案号}}, {{当事人名称}}, {{案件名称}}")
    print("- {{个人姓名}}, {{性别}}, {{身份证号}}")
    print("- {{案发地点}}, {{案件来源}}, {{案由}}")
    print("- {{违法事实}}, {{违法条款}}, {{处罚决定}}")
    print("- {{承办人}}, {{承办人姓名1}}, {{执法证号1}}")
    print("- {{项目名称}}, {{建设单位}}, {{施工单位}}")
    print("- {{生成日期}} (系统自动填充)")

if __name__ == '__main__':
    create_template()
